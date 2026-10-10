"""Analisa balancetes exportados do Dominio (um arquivo por empresa) e gera:

  - Resumo: contas com saldo e totais por empresa;
  - Zeramento: lancamentos sugeridos para zerar as contas escolhidas;
  - Mapeamento: contas lado a lado entre as empresas (de/para).

NAO altera nada no Dominio. Gera apenas uma planilha para voce conferir.

Uso:
    python zeramento_dominio.py PASTA_COM_BALANCETES --contrapartida 2.3.01.001
"""
import argparse
import csv
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook

# Nomes de coluna aceitos (minusculo, sem acento). Ajuste se o seu export for diferente.
ALIASES = {
    "conta": {"conta", "codigo", "cod", "codigo da conta", "reduzido"},
    "classificacao": {"classificacao", "classif", "classificacao da conta"},
    "descricao": {"descricao", "nome", "nome da conta", "historico"},
    "saldo": {"saldo", "saldo atual", "saldo final", "valor"},
    "natureza": {"natureza", "dc", "d/c", "tipo saldo"},
}
ACENTOS = str.maketrans("áàâãéêíóôõúç", "aaaaeeiooouc")


def norm(txt):
    return str(txt or "").strip().lower().translate(ACENTOS)


def para_numero(v):
    """Converte '1.234,56', '1234.56', '(10,00)' ou numero em float."""
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v or "").strip()
    if not s:
        return 0.0
    neg = s.startswith("(") and s.endswith(")") or s.endswith("-") or s.startswith("-")
    s = s.strip("()-").replace("R$", "").replace(" ", "")
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    try:
        n = float(s)
    except ValueError:
        return 0.0
    return -n if neg else n


def ler_linhas(arq):
    if arq.suffix.lower() == ".csv":
        with open(arq, newline="", encoding="utf-8-sig") as f:
            amostra = f.read(2048)
            f.seek(0)
            delim = ";" if amostra.count(";") > amostra.count(",") else ","
            return [list(r) for r in csv.reader(f, delimiter=delim)]
    ws = load_workbook(arq, data_only=True, read_only=True).active
    return [list(r) for r in ws.iter_rows(values_only=True)]


def ler_balancete(arq):
    """Retorna lista de contas: dict(conta, classificacao, descricao, saldo)
    onde saldo > 0 = devedor e saldo < 0 = credor."""
    linhas = ler_linhas(arq)
    cab = None
    for i, l in enumerate(linhas):
        nomes = [norm(c) for c in l]
        achou = {k: next((j for j, n in enumerate(nomes) if n in v), None)
                 for k, v in ALIASES.items()}
        if achou["saldo"] is not None and (achou["conta"] is not None
                                           or achou["classificacao"] is not None):
            cab, idx = i, achou
            break
    if cab is None:
        raise ValueError("cabecalho nao encontrado (preciso de colunas conta/classificacao e saldo)")

    def pega(l, k):
        j = idx[k]
        return l[j] if j is not None and j < len(l) else None

    contas = []
    for l in linhas[cab + 1:]:
        saldo = para_numero(pega(l, "saldo"))
        conta = str(pega(l, "conta") or "").strip()
        classif = str(pega(l, "classificacao") or "").strip()
        if not conta and not classif:
            continue
        nat = norm(pega(l, "natureza"))[:1]
        if nat == "c" and saldo > 0:
            saldo = -saldo
        contas.append({"conta": conta, "classificacao": classif,
                       "descricao": str(pega(l, "descricao") or "").strip(),
                       "saldo": round(saldo, 2)})
    return contas


def eh_analitica(c, todas):
    """Conta analitica = nenhuma outra classificacao comeca com a dela + '.'."""
    k = c["classificacao"]
    if not k:
        return True
    return not any(o["classificacao"].startswith(k + ".") for o in todas)


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("pasta", help="pasta com os balancetes (.xlsx ou .csv), um por empresa")
    p.add_argument("--saida", default="analise_zeramento.xlsx")
    p.add_argument("--prefixos", default="3,4",
                   help="classificacoes a zerar, separadas por virgula (padrao 3,4 = resultado). CONFIRA no seu plano de contas")
    p.add_argument("--contrapartida", default="", help="conta que recebe a contrapartida (ex.: apuracao do resultado)")
    p.add_argument("--data", default="", help="data do lancamento (dd/mm/aaaa)")
    a = p.parse_args()

    pasta = Path(a.pasta)
    arquivos = sorted(f for f in pasta.iterdir() if f.suffix.lower() in (".xlsx", ".csv"))
    if not arquivos:
        sys.exit(f"Nenhum .xlsx/.csv em {pasta}")
    prefixos = tuple(x.strip() for x in a.prefixos.split(",") if x.strip())

    wb = Workbook()
    ws_res, ws_zer, ws_map = wb.active, wb.create_sheet("Zeramento"), wb.create_sheet("Mapeamento")
    ws_res.title = "Resumo"
    ws_res.append(["Empresa", "Contas lidas", "Contas a zerar", "Total devedor", "Total credor", "Diferenca", "Observacao"])
    ws_zer.append(["Empresa", "Data", "Conta debito", "Conta credito", "Valor", "Historico", "Conta zerada", "Descricao"])
    mapa = {}  # classificacao -> {empresa: descricao}
    empresas = []

    for arq in arquivos:
        emp = arq.stem
        try:
            contas = ler_balancete(arq)
        except Exception as e:
            ws_res.append([emp, 0, 0, 0, 0, 0, f"ERRO: {e}"])
            print(f"[ERRO] {emp}: {e}")
            continue
        empresas.append(emp)
        alvo = [c for c in contas if c["saldo"] and c["classificacao"].startswith(prefixos)
                and eh_analitica(c, contas)]
        dev = sum(c["saldo"] for c in alvo if c["saldo"] > 0)
        cre = -sum(c["saldo"] for c in alvo if c["saldo"] < 0)
        obs = "" if a.contrapartida else "sem --contrapartida: preencha a conta na planilha"
        ws_res.append([emp, len(contas), len(alvo), round(dev, 2), round(cre, 2), round(dev - cre, 2), obs])
        for c in alvo:
            v = abs(c["saldo"])
            hist = f"Zeramento da conta {c['classificacao'] or c['conta']}"
            if c["saldo"] > 0:   # devedora: credita a conta para zerar
                deb, cred = a.contrapartida, c["conta"]
            else:                # credora: debita a conta para zerar
                deb, cred = c["conta"], a.contrapartida
            ws_zer.append([emp, a.data, deb, cred, v, hist, c["classificacao"], c["descricao"]])
        for c in contas:
            if c["classificacao"]:
                mapa.setdefault(c["classificacao"], {})[emp] = f"{c['conta']} - {c['descricao']}"
        print(f"[OK] {emp}: {len(contas)} contas, {len(alvo)} a zerar")

    ws_map.append(["Classificacao"] + empresas + ["Situacao"])
    for k in sorted(mapa):
        d = mapa[k]
        nomes = {x.split(" - ", 1)[-1].lower() for x in d.values()}
        sit = ("so em " + list(d)[0]) if len(d) == 1 and len(empresas) > 1 else (
            "descricao diferente" if len(nomes) > 1 else "igual")
        ws_map.append([k] + [d.get(e, "") for e in empresas] + [sit])

    for ws in wb:
        ws.freeze_panes = "A2"
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(
                50, max(12, max(len(str(c.value or "")) for c in col) + 2))
    wb.save(a.saida)
    print(f"\nPlanilha gerada: {a.saida}")


if __name__ == "__main__":
    main()
